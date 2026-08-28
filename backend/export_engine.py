"""Italian Professional Accounting & Legal Export Engine.
Validators (Codice Fiscale / Partita IVA), FatturaPA XML (practical v1.2, unsigned),
and pre-mapped accounting templates (XLSX/CSV) for IT accounting software.
"""
import io
import re
import csv
from datetime import datetime

# ----------------------------- Validators -----------------------------

_CF_ODD = {
    '0': 1, '1': 0, '2': 5, '3': 7, '4': 9, '5': 13, '6': 15, '7': 17, '8': 19, '9': 21,
    'A': 1, 'B': 0, 'C': 5, 'D': 7, 'E': 9, 'F': 13, 'G': 15, 'H': 17, 'I': 19, 'J': 21,
    'K': 2, 'L': 4, 'M': 18, 'N': 20, 'O': 11, 'P': 3, 'Q': 6, 'R': 8, 'S': 12, 'T': 14,
    'U': 16, 'V': 10, 'W': 22, 'X': 25, 'Y': 24, 'Z': 23,
}
_CF_EVEN = {
    '0': 0, '1': 1, '2': 2, '3': 3, '4': 4, '5': 5, '6': 6, '7': 7, '8': 8, '9': 9,
    'A': 0, 'B': 1, 'C': 2, 'D': 3, 'E': 4, 'F': 5, 'G': 6, 'H': 7, 'I': 8, 'J': 9,
    'K': 10, 'L': 11, 'M': 12, 'N': 13, 'O': 14, 'P': 15, 'Q': 16, 'R': 17, 'S': 18,
    'T': 19, 'U': 20, 'V': 21, 'W': 22, 'X': 23, 'Y': 24, 'Z': 25,
}


def valid_partita_iva(piva: str) -> bool:
    piva = re.sub(r"\D", "", piva or "")
    if len(piva) != 11:
        return False
    total = 0
    for i in range(11):
        d = int(piva[i])
        if i % 2 == 0:
            total += d
        else:
            d *= 2
            if d > 9:
                d -= 9
            total += d
    return total % 10 == 0


def valid_codice_fiscale(cf: str) -> bool:
    cf = (cf or "").strip().upper()
    # 16-char personal CF; 11-digit numeric CF (companies) validated as partita iva
    if re.fullmatch(r"\d{11}", cf):
        return valid_partita_iva(cf)
    if not re.fullmatch(r"[A-Z0-9]{16}", cf):
        return False
    if not re.fullmatch(r"[A-Z]{6}\d{2}[A-Z]\d{2}[A-Z]\d{3}[A-Z]", cf):
        return False
    total = 0
    for i in range(15):
        ch = cf[i]
        total += _CF_ODD[ch] if (i % 2 == 0) else _CF_EVEN[ch]
    return chr(ord('A') + (total % 26)) == cf[15]


def validate_field(kind: str, value: str):
    kind = (kind or "").lower()
    if kind in ("piva", "partita_iva", "p.iva"):
        return {"valid": valid_partita_iva(value), "kind": "piva"}
    if kind in ("cf", "codice_fiscale"):
        return {"valid": valid_codice_fiscale(value), "kind": "cf"}
    return {"valid": True, "kind": kind}


# ----------------------------- Helpers -----------------------------

def _num(s):
    """Parse an Italian-formatted number string to float. '1.220,00' -> 1220.00."""
    if s is None:
        return 0.0
    if isinstance(s, (int, float)):
        return float(s)
    s = re.sub(r"[^\d,.\-]", "", str(s))
    if not s:
        return 0.0
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _amt(s):
    return f"{_num(s):.2f}"


def _norm_date(s):
    """Return ISO yyyy-mm-dd from dd/mm/yyyy if possible, else original."""
    s = (s or "").strip()
    m = re.match(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})", s)
    if m:
        d, mo, y = m.groups()
        if len(y) == 2:
            y = "20" + y
        try:
            return f"{int(y):04d}-{int(mo):02d}-{int(d):02d}"
        except ValueError:
            return s
    return s


# ----------------------------- FatturaPA XML -----------------------------

_FPA_NS = "http://ivaservizi.agenziaentrate.gov.it/docs/xsd/fatture/v1.2"


def build_fatturapa_xml(structured: dict) -> bytes:
    """Practical FatturaPA v1.2 XML (unsigned, draft). Not SDI-signed."""
    from lxml import etree

    s = structured or {}
    sup = s.get("supplier", {}) or {}
    cus = s.get("customer", {}) or {}
    doc = s.get("document", {}) or {}
    tot = s.get("totals", {}) or {}
    vat_lines = s.get("vat_lines", []) or []
    items = s.get("line_items", []) or []
    iban = (s.get("iban") or "").strip()

    root = etree.Element("{%s}FatturaElettronica" % _FPA_NS, nsmap={"p": _FPA_NS})
    root.set("versione", "FPR12")

    def sub(parent, tag, text=None):
        e = etree.SubElement(parent, tag)
        if text is not None and text != "":
            e.text = str(text)
        return e

    # ---- Header ----
    header = sub(root, "FatturaElettronicaHeader")
    dt = sub(header, "DatiTrasmissione")
    idt = sub(dt, "IdTrasmittente")
    sub(idt, "IdPaese", "IT")
    sub(idt, "IdCodice", re.sub(r"\D", "", sup.get("piva", "")) or "00000000000")
    sub(dt, "ProgressivoInvio", "00001")
    sub(dt, "FormatoTrasmissione", "FPR12")
    sub(dt, "CodiceDestinatario", "0000000")

    cedente = sub(header, "CedentePrestatore")
    da = sub(cedente, "DatiAnagrafici")
    if sup.get("piva"):
        idiva = sub(da, "IdFiscaleIVA")
        sub(idiva, "IdPaese", "IT")
        sub(idiva, "IdCodice", re.sub(r"\D", "", sup.get("piva", "")))
    if sup.get("codice_fiscale"):
        sub(da, "CodiceFiscale", sup.get("codice_fiscale"))
    anag = sub(da, "Anagrafica")
    sub(anag, "Denominazione", sup.get("name") or "N/D")
    sub(da, "RegimeFiscale", "RF01")
    sede = sub(cedente, "Sede")
    sub(sede, "Indirizzo", sup.get("address") or "N/D")
    sub(sede, "CAP", "00000")
    sub(sede, "Comune", "N/D")
    sub(sede, "Nazione", "IT")

    committente = sub(header, "CessionarioCommittente")
    da2 = sub(committente, "DatiAnagrafici")
    if cus.get("piva"):
        idiva2 = sub(da2, "IdFiscaleIVA")
        sub(idiva2, "IdPaese", "IT")
        sub(idiva2, "IdCodice", re.sub(r"\D", "", cus.get("piva", "")))
    if cus.get("codice_fiscale"):
        sub(da2, "CodiceFiscale", cus.get("codice_fiscale"))
    anag2 = sub(da2, "Anagrafica")
    sub(anag2, "Denominazione", cus.get("name") or "N/D")
    sede2 = sub(committente, "Sede")
    sub(sede2, "Indirizzo", cus.get("address") or "N/D")
    sub(sede2, "CAP", "00000")
    sub(sede2, "Comune", "N/D")
    sub(sede2, "Nazione", "IT")

    # ---- Body ----
    body = sub(root, "FatturaElettronicaBody")
    dg = sub(body, "DatiGenerali")
    dgd = sub(dg, "DatiGeneraliDocumento")
    sub(dgd, "TipoDocumento", "TD01")
    sub(dgd, "Divisa", tot.get("currency") or "EUR")
    sub(dgd, "Data", _norm_date(doc.get("date", "")) or datetime.now().strftime("%Y-%m-%d"))
    sub(dgd, "Numero", doc.get("number") or "1")
    if tot.get("totale"):
        sub(dgd, "ImportoTotaleDocumento", _amt(tot.get("totale")))

    dbs = sub(body, "DatiBeniServizi")
    if items:
        for i, it in enumerate(items, 1):
            dl = sub(dbs, "DettaglioLinee")
            sub(dl, "NumeroLinea", str(i))
            sub(dl, "Descrizione", it.get("description") or "N/D")
            if it.get("quantity"):
                sub(dl, "Quantita", _amt(it.get("quantity")))
            if it.get("unit_price"):
                sub(dl, "PrezzoUnitario", _amt(it.get("unit_price")))
            sub(dl, "PrezzoTotale", _amt(it.get("total") or it.get("unit_price") or 0))
            sub(dl, "AliquotaIVA", _amt(it.get("vat_rate") or 0))
    else:
        dl = sub(dbs, "DettaglioLinee")
        sub(dl, "NumeroLinea", "1")
        sub(dl, "Descrizione", "Riepilogo documento")
        sub(dl, "PrezzoTotale", _amt(tot.get("imponibile") or 0))
        sub(dl, "AliquotaIVA", _amt(vat_lines[0].get("aliquota") if vat_lines else 0))

    if vat_lines:
        for vl in vat_lines:
            dr = sub(dbs, "DatiRiepilogo")
            sub(dr, "AliquotaIVA", _amt(vl.get("aliquota") or 0))
            sub(dr, "ImponibileImporto", _amt(vl.get("imponibile") or 0))
            sub(dr, "Imposta", _amt(vl.get("imposta") or 0))
    else:
        dr = sub(dbs, "DatiRiepilogo")
        sub(dr, "AliquotaIVA", "22.00")
        sub(dr, "ImponibileImporto", _amt(tot.get("imponibile") or 0))
        sub(dr, "Imposta", _amt(tot.get("imposta") or 0))

    if iban:
        dp = sub(body, "DatiPagamento")
        sub(dp, "CondizioniPagamento", "TP02")
        ddp = sub(dp, "DettaglioPagamento")
        sub(ddp, "ModalitaPagamento", "MP05")
        sub(ddp, "ImportoPagamento", _amt(tot.get("totale") or 0))
        sub(ddp, "IBAN", re.sub(r"\s", "", iban))

    return etree.tostring(root, pretty_print=True, xml_declaration=True, encoding="UTF-8")


# ----------------------------- Accounting templates -----------------------------

SOFTWARE_COLUMNS = {
    "zucchetti": ["Data Registrazione", "Tipo Doc", "Numero Doc", "P.IVA", "Codice Fiscale",
                  "Denominazione", "Imponibile", "Aliquota IVA", "Imposta", "Totale",
                  "Codice Causale", "Conto Contabile"],
    "teamsystem": ["DataReg", "TipoDoc", "NumeroDoc", "PartitaIVA", "CodiceFiscale",
                   "RagioneSociale", "Imponibile", "CodiceIVA", "Imposta", "Totale",
                   "Sezionale", "ContoContabile"],
    "datev": ["Datum", "Belegart", "Belegnummer", "USt-IdNr", "SteuerNr", "Name",
              "Netto", "Steuersatz", "Steuer", "Brutto", "Konto", "Gegenkonto"],
    "passepartout": ["Data", "Tipo Documento", "Protocollo", "Partita IVA", "Codice Fiscale",
                     "Cliente/Fornitore", "Imponibile", "% IVA", "Imposta", "Totale",
                     "Causale Contabile", "Conto"],
}

SOFTWARE_LABELS = {
    "zucchetti": "Zucchetti", "teamsystem": "TeamSystem",
    "datev": "Datev Koinos", "passepartout": "Passepartout",
}

DISCLAIMER = ("Tracciato generico da adattare - NON ufficiale. "
              "Verificare e mappare le colonne secondo la versione del gestionale in uso.")


def _invoice_rows(structured: dict, columns):
    """Best-effort mapping of invoice structured data onto a software column set."""
    s = structured or {}
    sup = s.get("supplier", {}) or {}
    doc = s.get("document", {}) or {}
    tot = s.get("totals", {}) or {}
    vat_lines = s.get("vat_lines", []) or []
    aliquota = vat_lines[0].get("aliquota") if vat_lines else ""
    imposta = vat_lines[0].get("imposta") if vat_lines else tot.get("imposta", "")

    base = {
        "data": doc.get("date", ""), "datum": doc.get("date", ""),
        "tipo": doc.get("type", "Fattura"), "numero": doc.get("number", ""),
        "protocollo": doc.get("number", ""), "belegnummer": doc.get("number", ""),
        "piva": sup.get("piva", ""), "cf": sup.get("codice_fiscale", ""),
        "name": sup.get("name", ""),
        "imponibile": tot.get("imponibile", ""), "aliquota": aliquota,
        "imposta": imposta, "totale": tot.get("totale", ""),
    }

    def cell(col):
        c = col.lower()
        if "data" in c or "datum" in c:
            return base["data"]
        if "belegart" in c or c.startswith("tipo"):
            return base["tipo"]
        if "numero" in c or "protocollo" in c or "belegnummer" in c:
            return base["numero"]
        if "ust-idnr" in c or "p.iva" in c or "partita" in c:
            return base["piva"]
        if "steuernr" in c or "fiscale" in c:
            return base["cf"]
        if "denominazione" in c or "ragionesociale" in c or "cliente" in c or c == "name":
            return base["name"]
        if "netto" in c or "imponibile" in c:
            return base["imponibile"]
        if "steuersatz" in c or "aliquota" in c or "% iva" in c or "codiceiva" in c:
            return base["aliquota"]
        if "steuer" == c or c == "imposta":
            return base["imposta"]
        if "brutto" in c or "totale" in c:
            return base["totale"]
        return ""

    return [[cell(col) for col in columns]]


def _flatten(structured, prefix=""):
    rows = []
    if isinstance(structured, dict):
        for k, v in structured.items():
            rows += _flatten(v, f"{prefix}{k}.")
    elif isinstance(structured, list):
        for i, v in enumerate(structured):
            rows += _flatten(v, f"{prefix}{i}.")
    else:
        rows.append((prefix.rstrip("."), "" if structured is None else str(structured)))
    return rows


def build_accounting_workbook(structured, category, software, doc_type="") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    software = (software or "zucchetti").lower()
    columns = SOFTWARE_COLUMNS.get(software, SOFTWARE_COLUMNS["zucchetti"])
    label = SOFTWARE_LABELS.get(software, software.title())

    wb = Workbook()
    ws = wb.active
    ws.title = f"Tracciato {label}"[:31]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(columns), 4))
    note = ws.cell(row=1, column=1, value=DISCLAIMER)
    note.font = Font(italic=True, color="B45309", size=9)
    note.alignment = Alignment(wrap_text=True)

    head_fill = PatternFill("solid", fgColor="6B5CE7")
    for j, col in enumerate(columns, 1):
        c = ws.cell(row=3, column=j, value=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        ws.column_dimensions[c.column_letter].width = max(12, len(col) + 2)

    rows = _invoice_rows(structured, columns) if category == "invoice" else []
    for i, row in enumerate(rows, 4):
        for j, val in enumerate(row, 1):
            ws.cell(row=i, column=j, value=val)

    # Raw data sheet
    ws2 = wb.create_sheet("Dati estratti")
    ws2.cell(row=1, column=1, value="Campo").font = Font(bold=True)
    ws2.cell(row=1, column=2, value="Valore").font = Font(bold=True)
    for i, (k, v) in enumerate(_flatten(structured), 2):
        ws2.cell(row=i, column=1, value=k)
        ws2.cell(row=i, column=2, value=v)
    ws2.column_dimensions["A"].width = 34
    ws2.column_dimensions["B"].width = 44

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_accounting_csv(structured, category, software) -> bytes:
    software = (software or "zucchetti").lower()
    columns = SOFTWARE_COLUMNS.get(software, SOFTWARE_COLUMNS["zucchetti"])
    rows = _invoice_rows(structured, columns) if category == "invoice" else []
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([f"# {DISCLAIMER}"])
    w.writerow(columns)
    for row in rows:
        w.writerow(row)
    return buf.getvalue().encode("utf-8-sig")
